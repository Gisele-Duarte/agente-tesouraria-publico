"""Gera o livro caixa em Excel, com fórmulas vivas (a planilha recalcula se alguém editar).

Mesmo layout da REV02: Instruções, Resumo Geral, Conferência Livro, Pendências,
uma aba por mês, Modelo, Extrato Bancário e Comprovantes WhatsApp.
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .modelo import FORMAS, STATUS, TIPOS, CONFIRMADO, PENDENTE, DIVERGENCIA, nome_aba, nome_mes_extenso, reais
from .processar import Livro
from .regras import PREFIXO_PREBENDA, PREFIXO_REPASSE_SEDE, PREFIXO_TRANSFERENCIA, eh_transferencia, motivo_na_sede

AZUL_ESCURO = PatternFill("solid", fgColor="1F3864")
AMARELO = PatternFill("solid", fgColor="FFF2CC")
CINZA = PatternFill("solid", fgColor="D9E1F2")
CORES_STATUS = {CONFIRMADO: "E2EFDA", PENDENTE: "FFF2CC", DIVERGENCIA: "F8CBAD"}
# Linhas novas e alteradas em relação à planilha base (mesmo código de cores da REV03/REV04).
CORES_ALTERACAO = {"Novo": PatternFill("solid", fgColor="DDEBF7"), "Alterado": PatternFill("solid", fgColor="E4DFEC")}
FONTE_TITULO = Font(bold=True, size=14, color="1F3864")
FONTE_CABECALHO = Font(bold=True, color="FFFFFF")
FONTE_DIGITADO = Font(color="0000FF")      # azul = valor digitado
FONTE_VINCULO = Font(color="008000")       # verde = vínculo com outra aba
NEGRITO = Font(bold=True)
MOEDA = '"R$" #,##0.00;[Red]-"R$" #,##0.00'
DATA = "DD/MM/YYYY"
LINHAS_EM_BRANCO = 3   # linhas vazias no fim de cada mês para lançar à mão

# Mesmos prefixos de regras.py, com o curinga do SUMIFS (que já ignora maiúsculas).
SE_PREBENDA = PREFIXO_PREBENDA.capitalize() + "*"
SE_REPASSE = PREFIXO_REPASSE_SEDE.title() + "*"
SE_TRANSFERENCIA = PREFIXO_TRANSFERENCIA.capitalize() + "*"

ROTULO_SEDE = "Repasse Igreja Sede"   # a importação procura este rótulo para ler a porcentagem do mês
PREFIXO_ABONO_PARCIAL = "Abono parcial de"   # nota da coluna E, lida de volta pela importação
ROTULO_PREBENDA_NO_SALDO = "Prebenda no saldo"
ROTULO_ABONADO = "ABONADO"            # coluna G do repasse à sede: o valor da coluna F não entra no cálculo
ABA_GESTAO_ANTERIOR = "Gestão anterior"
ABA_NA_SEDE = "Valores na Sede"
# Transferências internas ficam fora das abas dos meses (pedido da tesouraria): não mudam entradas,
# despesas nem saldo, só mostram dinheiro trocando de lugar. A coluna "Mês" permite reimportar.
ABA_TRANSFERENCIAS = "Transferências internas"
# Entradas que a tesouraria disse ser cartão, sem comprovante: ficam no caixa da congregação até a sede confirmar.
ABA_CARTAO_A_CONFIRMAR = "Cartão a confirmar"
MARCA_CARTAO_A_CONFIRMAR = "CARTÃO A CONFIRMAR COM A SEDE"
CABECALHO_TRANSFERENCIAS = ["Data", "Entrada/Saída", "Descrição", "Valor", "Forma", "Comprovante",
                            "Conferência", "Observação", "Código", "Mês"]

# Células fixas da aba Instruções usadas pelas fórmulas.
CEL_SEDE, CEL_PREBENDA, CEL_SALDO_INICIAL = "$C$6", "$C$7", "$C$10"


@dataclass(frozen=True)
class PosicoesMes:
    """Linhas do bloco de resumo de uma aba mensal (para o Resumo Geral apontar)."""
    entradas: int

    @property
    def sede(self): return self.entradas + 1
    @property
    def prebenda(self): return self.entradas + 2
    @property
    def local(self): return self.entradas + 3
    @property
    def despesas(self): return self.entradas + 4
    @property
    def saldo_mes(self): return self.entradas + 5
    @property
    def saldo_anterior(self): return self.entradas + 6
    @property
    def saldo_total(self): return self.entradas + 7
    @property
    def prebenda_paga(self): return self.entradas + 10
    @property
    def sede_paga(self): return self.entradas + 11
    @property
    def confirmados(self): return self.entradas + 14
    @property
    def pendentes(self): return self.entradas + 15
    @property
    def divergencias(self): return self.entradas + 16
    @property
    def sem_valor(self): return self.entradas + 17


def _faixa(coluna: str, ultima: int, aba: str = "") -> str:
    prefixo = f"'{aba}'!" if aba else ""
    return f"{prefixo}${coluna}$5:${coluna}${ultima}"


def gerar(livro: Livro, destino: Path) -> Path:
    achado = re.search(r"REV0*(\d+)", destino.stem, re.IGNORECASE)
    aba_alteracoes = f"Alterações REV{int(achado.group(1)):02d}" if achado else "Alterações"
    wb = Workbook()
    wb.remove(wb.active)
    _instrucoes(wb.create_sheet("Instruções"), livro, aba_alteracoes)
    resumo = wb.create_sheet("Resumo Geral")
    conferencia = wb.create_sheet("Conferência Livro")
    _alteracoes(wb.create_sheet(aba_alteracoes), livro)
    _pendencias(wb.create_sheet("Pendências"), livro)
    if livro.gestao_anterior:
        _gestao_anterior(wb.create_sheet(ABA_GESTAO_ANTERIOR), livro.gestao_anterior)
    if livro.regras.na_sede_desde:
        _valores_na_sede(wb.create_sheet(ABA_NA_SEDE), livro)
    a_confirmar = [l for l in livro.lancamentos if MARCA_CARTAO_A_CONFIRMAR in l.observacao]
    if a_confirmar:
        _tabela(wb.create_sheet(ABA_CARTAO_A_CONFIRMAR), "CARTÃO A CONFIRMAR COM A SEDE",
                "Entradas que a tesouraria identificou como cartão, sem comprovante. Contam como dinheiro da "
                "congregação até o extrato da maquininha da sede confirmar; o que for confirmado passa para "
                "'Cartão' pelo respostas.yaml.",
                ["Data", "Mês", "Descrição", "Valor", "Código"],
                [[l.data, l.mes, l.descricao, l.valor, l.id] for l in a_confirmar], [12, 9, 50, 13, 9], moeda=[4])
        fim = 4 + len(a_confirmar)
        ws_total = wb[ABA_CARTAO_A_CONFIRMAR]
        ws_total.cell(fim + 2, 3, "TOTAL").font = NEGRITO
        cel = ws_total.cell(fim + 2, 4, f"=SUM(D5:D{fim})")
        cel.number_format, cel.font = MOEDA, NEGRITO
    transferencias = [l for l in livro.lancamentos if eh_transferencia(l)]
    if transferencias:
        _transferencias(wb.create_sheet(ABA_TRANSFERENCIAS), transferencias, livro.cor_por_id)
    posicoes: dict[str, tuple[PosicoesMes, int]] = {}
    anterior = None
    for mes in livro.meses:
        ws = wb.create_sheet(nome_aba(mes))
        lancs = [l for l in livro.lancamentos if l.mes == mes and not eh_transferencia(l)]
        posicoes[mes] = _aba_mes(ws, mes, lancs, livro.fontes.get(mes, ""), anterior, livro.cor_por_id,
                                 livro.regras.abono_do_mes(mes), livro.regras.prebenda_quando_paga)
        anterior = (nome_aba(mes), posicoes[mes][0])
    _aba_mes(wb.create_sheet("Modelo"), None, [], "Copie esta aba para um mês novo.", None, {}, None,
             livro.regras.prebenda_quando_paga)
    _resumo(resumo, livro, posicoes)
    _conferencia(conferencia, livro, posicoes)
    _tabela(wb.create_sheet("Extrato Bancário"), livro.extrato_titulo or "EXTRATOS BANCÁRIOS",
            livro.extrato_periodos, ["Data", "Histórico (extrato)", "Valor", "Situação na planilha", "Arquivo"],
            [[l["data"], l["historico"], l["valor"], l.get("situacao"), l.get("arquivo")] for l in livro.extrato],
            [12, 48, 14, 50, 50], moeda=[3])
    _tabela(wb.create_sheet("Comprovantes WhatsApp"), "ÍNDICE DE ANEXOS DO WHATSAPP",
            '"Pendente – Etapa 3" = foto ou página da agenda que ainda precisa de leitura.',
            ["Data", "Enviado por", "Tipo", "Valor lido", "Pagador", "Recebedor", "Mensagem/legenda",
             "Uso na planilha", "Arquivo", "ZIP"],
            [[c["data"], c.get("enviado_por"), c.get("tipo"), c.get("valor"), c.get("pagador"),
              c.get("recebedor"), c.get("mensagem"), c.get("uso"), c.get("arquivo"), c.get("zip")]
             for c in livro.comprovantes],
            [12, 14, 24, 13, 30, 30, 50, 30, 42, 26], moeda=[4])
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino


# --- abas ---------------------------------------------------------------
def _titulo(ws, texto: str, subtitulo: str = "") -> None:
    ws["A1"], ws["A1"].font = texto, FONTE_TITULO
    if subtitulo:
        ws["A2"] = subtitulo
        ws["A2"].alignment = Alignment(wrap_text=False)


def _cabecalho_celula(cel) -> None:
    cel.font, cel.fill = FONTE_CABECALHO, AZUL_ESCURO
    cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _cabecalho(ws, linha: int, titulos: list[str]) -> None:
    for coluna, texto in enumerate(titulos, start=1):
        cel = ws.cell(linha, coluna, texto)
        cel.font, cel.fill = FONTE_CABECALHO, AZUL_ESCURO
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _larguras(ws, larguras: list[float]) -> None:
    for i, largura in enumerate(larguras):
        ws.column_dimensions[chr(ord("A") + i)].width = largura


def _instrucoes(ws, livro: Livro, aba_alteracoes: str) -> None:
    ws["B1"], ws["B1"].font = livro.titulo or f"LIVRO CAIXA – {livro.nome}", FONTE_TITULO
    ws["B2"] = "Organização dos lançamentos: livro caixa + WhatsApp + extratos"
    ws["B4"], ws["B4"].font = "PARÂMETROS (células amarelas podem ser alteradas)", NEGRITO
    parametros = [
        (6, "Repasse Igreja Sede", float(livro.regras.repasse_sede), "Percentual sobre as entradas do mês"),
        (7, "Prebenda Pastoral", float(livro.regras.prebenda), "Percentual sobre as entradas do mês"),
        (8, "Saldo da Igreja Local", "=1-C6-C7", "Calculado: o que sobra das entradas para a igreja local"),
        (9, ROTULO_PREBENDA_NO_SALDO, "quando paga" if livro.regras.prebenda_quando_paga else "todo mês",
         "'quando paga': o saldo só é debitado pela prebenda efetivamente paga (adiantamentos); o restante dos "
         "20% fica sinalizado como 'a pagar'. 'todo mês': os 20% saem do saldo todo mês."),
        (10, "Saldo inicial antes do 1º lançamento",
         float(livro.saldo_inicial or 0),
         "Informado pela tesouraria" if livro.saldo_inicial is not None
         else "PREENCHER: saldo em caixa/banco antes do 1º lançamento (ainda não informado)"),
    ]
    for linha, rotulo, valor, nota in parametros:
        ws.cell(linha, 2, rotulo)
        cel = ws.cell(linha, 3, valor)
        cel.number_format = MOEDA if linha == 10 else "@" if linha == 9 else "0%"
        if linha != 8:
            cel.fill, cel.font = AMARELO, FONTE_DIGITADO
        ws.cell(linha, 4, nota)
    textos = [
        ("COMO O SALDO É CALCULADO", True),
        ('1. Entradas do mês = soma das linhas "Entrada" da aba do mês.', False),
        ("2. Repasse à sede (20%) e prebenda (20%) são calculados sobre as entradas; o restante (60%) fica com a "
         "igreja local. Nos meses de abono da sede, o repasse é 0% e a igreja local fica com 80%. A porcentagem "
         "da sede de cada mês fica na coluna D do resumo da aba do mês.", False),
        *[(f"   • Repasse à sede abonado de {nome_mes_extenso(a.de)} a {nome_mes_extenso(a.ate)}"
           + (f" (parcial: {reais(a.valor)})" if a.valor is not None else "")
           + (f": {a.motivo}" if a.motivo else ""), False) for a in livro.regras.abonos_sede],
        *([('3. Prebenda: os 20% são sinalizados como "a pagar" (Resumo Geral, coluna "Prebenda a pagar"). '
             'O saldo só é debitado nos meses em que houve pagamento/adiantamento de prebenda ("Prebenda…"), '
             'pelo valor pago.', False),
           ("4. Saldo total = entradas − repasse à sede − despesas − prebenda paga no mês + saldo do mês anterior. "
            'Repasses à sede ("Repasse Sede…") não são despesa: saem da parte da sede.', False)]
          if livro.regras.prebenda_quando_paga else
          [('3. Despesas do mês = saídas operacionais. Pagamentos de prebenda ("Prebenda…") e repasses à sede '
            '("Repasse Sede…") NÃO são descontados da parte da igreja local, porque já saíram das partes de 20%.',
            False),
           ("4. Saldo total = parte da igreja local − despesas do mês + saldo do mês anterior.", False)]),
        ('5. "Transferência interna…" (banco → caixa do pastor) não é entrada nem despesa.', False),
        ("", False),
        ("COMO ALIMENTAR", True),
        ("• Esta planilha é gerada pelo Agente de Tesouraria. Correções feitas aqui à mão não voltam para o "
         "agente: registre-as em respostas.yaml para a próxima revisão já sair certa.", False),
        ('• Na Descrição, comece pela categoria: "Dízimo – Nome", "Oferta – culto (Pix)"…', False),
        ('• Pagamento de prebenda começa com "Prebenda"; repasse à sede com "Repasse Sede".', False),
        ("• Use o filtro da coluna Conferência para ver só 🟡 ou 🔴.", False),
        ("", False),
        ("LEGENDA DE CONFERÊNCIA", True),
        ("🟢 CONFIRMADO — temos data + valor + identificação suficiente.", False),
        ("🟡 PENDENTE — houve a movimentação, mas falta informação (valor, leitura incerta, descrição).", False),
        ("🔴 DIVERGÊNCIA — não bate com extrato/comprovante ou parece haver duplicidade.", False),
        ("", False),
        ("SITUAÇÃO ATUAL", True),
        (livro.situacao, False),
        ("", False),
        ("Texto azul = valor digitado · preto = fórmula · verde = vínculo com outra aba · "
         "amarelo = célula para preencher.", False),
        (f"Linhas com fundo azul = novas; lilás = alteradas em relação à revisão anterior (detalhe e motivo de "
         f"cada uma na aba {aba_alteracoes}). A coluna Código identifica o lançamento nas respostas.", False),
    ]
    for i, (texto, negrito) in enumerate(textos, start=12):
        cel = ws.cell(i, 2, texto)
        cel.font = NEGRITO if negrito else Font()
        cel.alignment = Alignment(wrap_text=True, vertical="top")
    _larguras(ws, [3, 110, 14, 70])


def _aba_mes(ws, mes: str | None, lancs, fonte: str, anterior, cores: dict, abono,
             prebenda_quando_paga: bool) -> tuple[PosicoesMes, int]:
    titulo = f"MOVIMENTO DO CAIXA – {nome_mes_extenso(mes)}" if mes else "MOVIMENTO DO CAIXA – <mês>"
    _titulo(ws, titulo, fonte)
    _cabecalho(ws, 4, ["Data", "Entrada/Saída", "Descrição (Categoria – detalhe)", "Valor", "Forma",
                       "Comprovante", "Conferência", "Observação", "Código"])
    for i, l in enumerate(lancs, start=5):
        valores = [l.data, l.tipo, l.descricao, None if l.valor is None else float(l.valor), l.forma,
                   l.comprovante, l.status, l.observacao, l.id]
        for coluna, valor in enumerate(valores, start=1):
            cel = ws.cell(i, coluna, valor)
            cel.font = FONTE_DIGITADO
            if l.id in cores and coluna != 7:
                cel.fill = CORES_ALTERACAO[cores[l.id]]
        ws.cell(i, 1).number_format = DATA
        ws.cell(i, 4).number_format = MOEDA
    ultima = 4 + max(len(lancs) + LINHAS_EM_BRANCO, 60 if mes is None else 0)
    for linha in range(5, ultima + 1):
        ws.cell(linha, 1).number_format = DATA
        ws.cell(linha, 4).number_format = MOEDA
    d, b, c, g = (_faixa(coluna, ultima) for coluna in "DBCG")
    r = ultima + 2
    ws.cell(r, 2, "RESUMO DO MÊS").font = NEGRITO
    pos = PosicoesMes(r + 1)
    if anterior:
        saldo_anterior = f"='{anterior[0]}'!$C${anterior[1].saldo_total}"
    else:
        saldo_anterior = f"=Instruções!{CEL_SALDO_INICIAL}"
    linhas = [
        (pos.entradas, "Entradas (total)",
         f'=SUMIFS({d},{b},"Entrada")-SUMIFS({d},{b},"Entrada",{c},"{SE_TRANSFERENCIA}")'),
        (pos.sede, ROTULO_SEDE, f"=C{pos.entradas}*D{pos.sede}"
                                + (f"-{float(abono.valor)}" if abono and abono.valor is not None else "")),
        (pos.prebenda, "Prebenda Pastoral (a pagar)" if prebenda_quando_paga else "Prebenda Pastoral",
         f"=C{pos.entradas}*D{pos.prebenda}"),
        (pos.local, "Saldo da Igreja Local",
         f"=C{pos.entradas}-C{pos.sede}" + ("" if prebenda_quando_paga else f"-C{pos.prebenda}")),
        (pos.despesas, "(−) Despesas do mês",
         f'=SUMIFS({d},{b},"Saída")-SUMIFS({d},{b},"Saída",{c},"{SE_PREBENDA}")'
         f'-SUMIFS({d},{b},"Saída",{c},"{SE_REPASSE}")-SUMIFS({d},{b},"Saída",{c},"{SE_TRANSFERENCIA}")'),
        (pos.saldo_mes, "Saldo do mês (local − despesas − prebenda paga)" if prebenda_quando_paga
         else "Saldo do mês (local − despesas)",
         f"=C{pos.local}-C{pos.despesas}" + (f"-C{pos.prebenda_paga}" if prebenda_quando_paga else "")),
        (pos.saldo_anterior, "Saldo do mês anterior", saldo_anterior),
        (pos.saldo_total, "SALDO TOTAL", f"=C{pos.saldo_mes}+C{pos.saldo_anterior}"),
        (pos.prebenda_paga, "Prebenda paga no mês", f'=SUMIFS({d},{b},"Saída",{c},"{SE_PREBENDA}")'),
        (pos.sede_paga, "Repasse à sede pago no mês", f'=SUMIFS({d},{b},"Saída",{c},"{SE_REPASSE}")'),
        (pos.confirmados, "🟢 Confirmados", f'=COUNTIF({g},"{CONFIRMADO}")'),
        (pos.pendentes, "🟡 Pendentes", f'=COUNTIF({g},"{PENDENTE}")'),
        (pos.divergencias, "🔴 Divergências", f'=COUNTIF({g},"{DIVERGENCIA}")'),
        (pos.sem_valor, "Lançamentos sem valor", f'=COUNTIFS({b},"<>",{d},"")'),
    ]
    for linha, rotulo, formula in linhas:
        ws.cell(linha, 2, rotulo)
        cel = ws.cell(linha, 3, formula)
        if linha < pos.confirmados:
            cel.number_format = MOEDA
        if linha == pos.saldo_anterior and anterior:
            cel.font = FONTE_VINCULO
    # Porcentagens do mês (coluna D): a da sede é editável, 0% nos meses de abono.
    sede_pct = ws.cell(pos.sede, 4, 0 if abono and abono.valor is None else f"=Instruções!{CEL_SEDE}")
    sede_pct.fill, sede_pct.font = AMARELO, FONTE_DIGITADO
    ws.cell(pos.prebenda, 4, f"=Instruções!{CEL_PREBENDA}").font = FONTE_VINCULO
    ws.cell(pos.local, 4, f"=1-D{pos.sede}" + ("" if prebenda_quando_paga else f"-D{pos.prebenda}"))
    for linha in (pos.sede, pos.prebenda, pos.local):
        ws.cell(linha, 4).number_format = "0%"
    if abono and abono.valor is not None:
        ws.cell(pos.sede, 5, f"{PREFIXO_ABONO_PARCIAL} {reais(abono.valor)}" + (f": {abono.motivo}" if abono.motivo else ""))
    elif abono:
        ws.cell(pos.sede, 5, "Repasse à sede abonado neste mês" + (f": {abono.motivo}" if abono.motivo else ""))
    if abono:
        # Só informativo: quanto seriam os 20% (ou a parte abonada). Não entra em nenhuma fórmula do caixa.
        abonado = f"=C{pos.entradas}*Instruções!{CEL_SEDE}" if abono.valor is None else float(abono.valor)
        cel = ws.cell(pos.sede, 6, abonado)
        cel.number_format, cel.fill, cel.font = MOEDA, AMARELO, NEGRITO
        ws.cell(pos.sede, 7, ROTULO_ABONADO).font = Font(bold=True, color="C00000")
    if prebenda_quando_paga:
        ws.cell(pos.prebenda, 5, "Informativo: só sai do saldo a prebenda paga no mês (linha 'Prebenda paga no mês')")
    ws.cell(pos.despesas, 4, "Não inclui prebenda, repasse à sede nem transferências internas.")
    ws.cell(pos.saldo_total, 2).font = ws.cell(pos.saldo_total, 3).font = NEGRITO
    ws.cell(pos.prebenda_paga - 1, 2, "INFORMATIVO – pagamentos feitos com as partes de 20%").font = NEGRITO
    ws.cell(pos.confirmados - 1, 2, "CONFERÊNCIA DOS LANÇAMENTOS").font = NEGRITO

    for coluna, opcoes in (("B", TIPOS), ("E", FORMAS), ("G", STATUS)):
        validacao = DataValidation(type="list", formula1='"' + ",".join(opcoes) + '"', allow_blank=True)
        validacao.add(f"{coluna}5:{coluna}{ultima}")
        ws.add_data_validation(validacao)
    for status, cor in CORES_STATUS.items():
        ws.conditional_formatting.add(f"G5:G{ultima}", CellIsRule(operator="equal", formula=[f'"{status}"'],
                                                                   fill=PatternFill("solid", fgColor=cor)))
    ws.auto_filter.ref = f"A4:H{ultima}"
    ws.freeze_panes = "A5"
    _larguras(ws, [12, 14, 46, 14, 15, 22, 17, 60, 10])
    return pos, ultima


def _resumo(ws, livro: Livro, posicoes: dict) -> None:
    _titulo(ws, "RESUMO GERAL – mês a mês",
            'Tudo vinculado às abas mensais. "A pagar/repassar" = 20% devidos − valores já pagos '
            '(lançados como "Prebenda…" ou "Repasse Sede…").')
    _cabecalho(ws, 4, ["Mês", "Entradas", "Repasse Sede", "Prebenda (20%)", "Igreja Local", "Despesas",
                       "Saldo do mês", "Saldo anterior", "SALDO TOTAL", "Prebenda paga",
                       "Prebenda a pagar (acum.)", "Sede paga", "Sede a repassar (acum.)", "🟢", "🟡", "🔴"])
    linha = 4
    for linha, mes in enumerate(livro.meses, start=5):
        pos, _ = posicoes[mes]
        aba = nome_aba(mes)

        def ref(n: int) -> str:
            return f"='{aba}'!$C${n}"

        ws.cell(linha, 1, nome_mes_extenso(mes))
        valores = {2: ref(pos.entradas), 3: ref(pos.sede), 4: ref(pos.prebenda), 5: ref(pos.local),
                   6: ref(pos.despesas), 7: ref(pos.saldo_mes), 8: ref(pos.saldo_anterior),
                   9: ref(pos.saldo_total), 10: ref(pos.prebenda_paga),
                   11: f"=D{linha}-J{linha}" + (f"+K{linha - 1}" if linha > 5 else ""),
                   12: ref(pos.sede_paga),
                   13: f"=C{linha}-L{linha}" + (f"+M{linha - 1}" if linha > 5 else ""),
                   14: ref(pos.confirmados), 15: ref(pos.pendentes), 16: ref(pos.divergencias)}
        for coluna, formula in valores.items():
            cel = ws.cell(linha, coluna, formula)
            cel.font = FONTE_VINCULO if formula.startswith("='") else Font()
            if coluna < 14:
                cel.number_format = MOEDA
    total = linha + 1
    ws.cell(total, 1, "TOTAL").font = NEGRITO
    for coluna in range(2, 17):
        letra = chr(ord("A") + coluna - 1)
        formula = (f"={letra}{linha}" if letra in "IKM" else None if letra == "H"
                   else f"=SUM({letra}5:{letra}{linha})")
        if formula:
            cel = ws.cell(total, coluna, formula)
            cel.font, cel.fill = NEGRITO, CINZA
            if coluna < 14:
                cel.number_format = MOEDA
    if livro.regras.na_sede_desde:
        # Informativo: quanto do SALDO TOTAL está, na verdade, na conta da sede (não altera o saldo).
        for coluna, texto in ((17, "Na Sede (acum.)"), (18, "Saldo na congregação")):
            _cabecalho_celula(ws.cell(4, coluna, texto))
        for linha_mes, mes in enumerate(livro.meses, start=5):
            na_sede = (f"=IFERROR(INDEX('{ABA_NA_SEDE}'!$L:$L,MATCH(\"{mes}\",'{ABA_NA_SEDE}'!$J:$J,0)),"
                       + (f"Q{linha_mes - 1}" if linha_mes > 5 else "0") + ")")
            ws.cell(linha_mes, 17, na_sede).number_format = MOEDA
            ws.cell(linha_mes, 18, f"=I{linha_mes}-Q{linha_mes}").number_format = MOEDA
        for coluna, formula in ((17, f"=Q{linha}"), (18, f"=R{linha}")):
            cel = ws.cell(total, coluna, formula)
            cel.font, cel.fill, cel.number_format = NEGRITO, CINZA, MOEDA
    if livro.observacao_resumo:
        ws.cell(total + 2, 1, livro.observacao_resumo).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B5"
    _larguras(ws, [16] + [14] * 12 + [6, 6, 6] + ([15, 17] if livro.regras.na_sede_desde else []))


def _conferencia(ws, livro: Livro, posicoes: dict) -> None:
    _titulo(ws, "CONFERÊNCIA – soma dos lançamentos x totais escritos no livro",
            'Considera só as linhas vindas do livro (Comprovante começa com "Livro caixa"). '
            "Diferenças indicam leitura incerta ou soma errada no livro.")
    _cabecalho(ws, 4, ["Mês", "Entradas – soma dos lançamentos", "Entradas – total anotado no livro",
                       "Diferença", "Saídas – soma (sem prebenda)", "Saídas – total anotado no livro",
                       "Diferença", "Observação"])
    for linha, (mes, totais) in enumerate(sorted(livro.livro_totais.items()), start=5):
        if mes not in posicoes:
            continue
        aba, ultima = nome_aba(mes), posicoes[mes][1]
        d, b, c, f = (_faixa(coluna, ultima, aba) for coluna in "DBCF")
        ws.cell(linha, 1, nome_mes_extenso(mes))
        ws.cell(linha, 2, f'=SUMIFS({d},{b},"Entrada",{f},"Livro caixa*")')
        ws.cell(linha, 3, totais.get("entradas"))
        ws.cell(linha, 4, f'=IF(C{linha}="","",B{linha}-C{linha})')
        ws.cell(linha, 5, f'=SUMIFS({d},{b},"Saída",{f},"Livro caixa*")'
                          f'-SUMIFS({d},{b},"Saída",{f},"Livro caixa*",{c},"{SE_PREBENDA}")')
        ws.cell(linha, 6, totais.get("saidas"))
        ws.cell(linha, 7, f'=IF(F{linha}="","",E{linha}-F{linha})')
        ws.cell(linha, 8, totais.get("observacao", ""))
        for coluna in range(2, 8):
            ws.cell(linha, coluna).number_format = MOEDA
        for coluna in (3, 6):
            ws.cell(linha, coluna).font = FONTE_DIGITADO
    _larguras(ws, [16, 16, 16, 12, 16, 16, 12, 90])


def _alteracoes(ws, livro: Livro) -> None:
    contagem = {t: sum(a["tipo"] == t for a in livro.alteracoes) for t in ("Novo", "Alterado", "Removido")}
    _titulo(ws, "ALTERAÇÕES EM RELAÇÃO À REVISÃO ANTERIOR",
            f"{contagem['Novo']} lançamentos novos (azul), {contagem['Alterado']} alterados (lilás) e "
            f"{contagem['Removido']} removidos. 'Motivo' traz a resposta da tesouraria, a conferência feita ou a "
            "anotação da foto lida.")
    _cabecalho(ws, 4, ["Tipo", "Código", "Mês", "Data", "Entrada/Saída", "Descrição", "Valor", "O que mudou",
                       "Motivo"])
    for i, item in enumerate(livro.alteracoes, start=5):
        l = item["lanc"]
        valores = [item["tipo"], item["id"], nome_mes_extenso(l.mes), l.data, l.tipo, l.descricao,
                   None if l.valor is None else float(l.valor), item["mudou"], item["motivo"]]
        for coluna, valor in enumerate(valores, start=1):
            cel = ws.cell(i, coluna, valor)
            cel.alignment = Alignment(wrap_text=coluna >= 8, vertical="top")
            if item["tipo"] in CORES_ALTERACAO:
                cel.fill = CORES_ALTERACAO[item["tipo"]]
        ws.cell(i, 4).number_format = DATA
        ws.cell(i, 7).number_format = MOEDA
    ws.auto_filter.ref = f"A4:I{max(4 + len(livro.alteracoes), 5)}"
    ws.freeze_panes = "A5"
    _larguras(ws, [10, 10, 15, 11, 9, 46, 13, 50, 70])


def _pendencias(ws, livro: Livro) -> None:
    _titulo(ws, "PENDÊNCIAS PARA FECHAR A ORGANIZAÇÃO",
            "Cada item tem um código. Para resolver, responda citando o código: a resposta vai para "
            "respostas.yaml e entra na próxima revisão.")
    linha = 4

    def secao(texto: str) -> None:
        nonlocal linha
        linha += 1
        cel = ws.cell(linha, 1, texto)
        cel.font, cel.fill = NEGRITO, CINZA
        linha += 1

    def item(*valores) -> None:
        nonlocal linha
        for coluna, valor in enumerate(valores, start=1):
            ws.cell(linha, coluna, valor).alignment = Alignment(wrap_text=True, vertical="top")
        linha += 1

    # Seções 1–3 saem dos lançamentos e das fotos; a importação as ignora (ver SECOES_DERIVADAS).
    vermelhas = [p for p in livro.pendencias if p.cor == "🔴"]
    amarelas = [p for p in livro.pendencias if p.cor == "🟡"]
    secao(f"1. DIVERGÊNCIAS 🔴 ({len(vermelhas)})")
    for p in vermelhas:
        item(p.id, p.titulo, p.detalhe)
    secao(f"2. PENDENTES 🟡 ({len(amarelas)})")
    for p in amarelas:
        item(p.id, p.titulo, p.detalhe)
    secao(f"3. FOTOS DA AGENDA AINDA NÃO LIDAS ({len(livro.fotos_pendentes)})")
    for mes, texto in livro.fotos_pendentes_por_mes:
        item(mes, texto)
    numero = 4
    secao(f"{numero}. PERGUNTAS PARA A LIDERANÇA (coluna C = resposta registrada)")
    for id_, pergunta in livro.perguntas.items():
        item(pergunta["titulo"], pergunta["texto"], livro.perguntas_respondidas.get(id_, ""))
    for nome, avisos in _avisos_por_secao(livro.avisos).items():
        numero += 1
        secao(f"{numero}. {nome}")
        for aviso in avisos:
            item(aviso["titulo"], aviso["texto"])
    _larguras(ws, [22, 70, 90])


def _avisos_por_secao(avisos: list[dict]) -> dict[str, list[dict]]:
    grupos: dict[str, list[dict]] = {}
    for aviso in avisos:
        nome = re.sub(r"^\d+\.\s*", "", aviso.get("secao") or "AVISOS")
        grupos.setdefault(nome, [])
        if aviso.get("titulo") or aviso.get("texto"):
            grupos[nome].append({"titulo": aviso.get("titulo") or "—", "texto": aviso.get("texto", "")})
    return grupos


def _valores_na_sede(ws, livro: Livro) -> None:
    """Entradas do livro cujo dinheiro ficou na conta da sede: lista para conferir com a sede, totais por mês."""
    itens = [(l, m) for l in livro.lancamentos if (m := motivo_na_sede(l, livro.regras))]
    desde = nome_mes_extenso(livro.regras.na_sede_desde)
    _tabela(ws, "VALORES NA SEDE – entradas do livro que não caíram na conta da congregação",
            f"A partir de {desde}. " + (livro.regras.na_sede_motivo or "") +
            " Confira com a sede e corrija pelo respostas.yaml o que tiver caído em outro lugar.",
            ["Data", "Mês", "Descrição", "Valor", "Forma", "Motivo", "Comprovante", "Código"],
            [[l.data, l.mes, l.descricao, l.valor, l.forma, m, l.comprovante, l.id] for l, m in itens],
            [12, 9, 50, 13, 12, 42, 40, 9], moeda=[4])
    fim = 4 + max(len(itens), 1)
    col = 10
    ws.cell(4, col, "Mês").font = ws.cell(4, col + 1, "Na sede no mês").font = NEGRITO
    ws.cell(4, col + 2, "Acumulado").font = NEGRITO
    linha = 5
    for mes in livro.meses:
        if mes < livro.regras.na_sede_desde:
            continue
        ws.cell(linha, col, mes)
        ws.cell(linha, col + 1, f'=SUMIFS($D$5:$D${fim},$B$5:$B${fim},J{linha})').number_format = MOEDA
        ws.cell(linha, col + 2, f"=K{linha}" + (f"+L{linha - 1}" if linha > 5 else "")).number_format = MOEDA
        linha += 1
    ws.cell(linha, col, "TOTAL").font = NEGRITO
    tot = ws.cell(linha, col + 1, f"=SUM(K5:K{linha - 1})")
    tot.number_format, tot.font = MOEDA, NEGRITO
    _larguras(ws, [12, 9, 50, 13, 12, 42, 40, 9, 3, 10, 15, 15])


def _transferencias(ws, lancs, cores: dict) -> None:
    """Transferências internas (banco ↔ caixa do pastor, PicPay…), fora das abas dos meses."""
    _tabela(ws, "TRANSFERÊNCIAS INTERNAS – dinheiro que só trocou de lugar",
            "Não entram nas entradas nem nas despesas e não mudam o saldo. Ficam nesta aba, fora dos meses, "
            "para não confundir a leitura do caixa.",
            CABECALHO_TRANSFERENCIAS,
            [[l.data, l.tipo, l.descricao, l.valor, l.forma, l.comprovante, l.status, l.observacao, l.id, l.mes]
             for l in lancs],
            [12, 12, 60, 14, 14, 40, 16, 50, 9, 9], moeda=[4])
    for i, l in enumerate(lancs, start=5):
        if l.id in cores:
            for coluna in range(1, 11):
                if coluna != 7:
                    ws.cell(i, coluna).fill = CORES_ALTERACAO[cores[l.id]]
    fim = 4 + max(len(lancs), 1)
    for linha, (rotulo, tipo) in enumerate((("Total de saídas", "Saída"), ("Total de entradas", "Entrada")),
                                           start=fim + 2):
        ws.cell(linha, 3, rotulo).font = NEGRITO
        cel = ws.cell(linha, 4, f'=SUMIFS($D$5:$D${fim},$B$5:$B${fim},"{tipo}")')
        cel.number_format, cel.font = MOEDA, NEGRITO


def _gestao_anterior(ws, dados: dict) -> None:
    """Período anterior à tesouraria atual: só registro, não entra em nenhum saldo do livro."""
    linhas = [[l.get("data"), l.get("tipo"), l.get("descricao"), l.get("valor"), l.get("fonte"), l.get("observacao")]
              for l in dados.get("lancamentos") or []]
    _tabela(ws, dados.get("titulo") or "GESTÃO ANTERIOR",
            dados.get("subtitulo") or "Registro à parte: estes valores não entram no livro caixa.",
            ["Data", "Entrada/Saída", "Descrição", "Valor", "Fonte", "Observação"],
            linhas, [12, 13, 55, 14, 28, 55], moeda=[4])
    linha = 4 + len(linhas) + 2
    for item in dados.get("resumo") or []:
        ws.cell(linha, 3, item.get("rotulo")).font = NEGRITO
        valor = ws.cell(linha, 4, item.get("valor"))
        valor.number_format = MOEDA
        ws.cell(linha, 6, item.get("observacao"))
        linha += 1


def _tabela(ws, titulo: str, subtitulo: str, cabecalho: list[str], linhas: list[list],
            larguras: list[float], moeda: list[int]) -> None:
    _titulo(ws, titulo, subtitulo)
    _cabecalho(ws, 4, cabecalho)
    for i, valores in enumerate(linhas, start=5):
        for coluna, valor in enumerate(valores, start=1):
            if hasattr(valor, "is_finite"):  # Decimal
                valor = float(valor)
            ws.cell(i, coluna, valor)
        ws.cell(i, 1).number_format = DATA
        for coluna in moeda:
            ws.cell(i, coluna).number_format = MOEDA
    ws.auto_filter.ref = f"A4:{chr(ord('A') + len(cabecalho) - 1)}{max(4 + len(linhas), 5)}"
    ws.freeze_panes = "A5"
    _larguras(ws, larguras)
