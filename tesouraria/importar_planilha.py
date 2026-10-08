"""Importa uma revisão da planilha (REV02, REV04…) para os dados editáveis da congregação.

A planilha gerada pelo agente é a forma de ler; os arquivos em dados/ são a forma de guardar.
Importar uma revisão nova substitui os dados atuais, mas antes copia tudo para
cache/backup_dados_<data-hora>/ — nada é apagado.
"""
from __future__ import annotations

import re
import shutil
from datetime import datetime
from pathlib import Path

import openpyxl

from .dados import Congregacao, CAMPOS_COMPROVANTE
from .modelo import (TIPOS, MESES_PT, Lancamento, dinheiro, para_data, normalizar_status,
                     mes_da_aba, proximo_mes)

FIM_DOS_LANCAMENTOS = "RESUMO DO MÊS"
ABA_TRANSFERENCIAS = "Transferências internas"  # mesmo nome de planilha.ABA_TRANSFERENCIAS
SECOES_DERIVADAS = ("AGENDA", "FOTOS", "DIVERGÊNCIAS", "PENDENTES 🟡")


class ErroImportacao(Exception):
    pass


def importar(cong: Congregacao, planilha: Path, substituir: bool = False) -> dict:
    if cong.dados("lancamentos.csv").exists() and not substituir:
        raise ErroImportacao(
            "esta congregação já tem dados importados. Use --substituir para trocar pela planilha "
            "nova (os dados atuais vão para cache/ antes, nada é apagado).")
    wb = openpyxl.load_workbook(planilha, data_only=True)
    avisos: list[str] = []

    lancamentos, fontes = _lancamentos(wb, avisos)
    if not lancamentos:
        raise ErroImportacao("nenhuma aba mensal (ex.: Jun-2025) com lançamentos foi encontrada")
    config = _config(wb, planilha, cong.config())
    notas = {
        "titulo": config.get("titulo", ""),
        "situacao": _situacao(wb),
        "meses": fontes,
        "livro_totais": _livro_totais(wb, avisos),
        "observacao_resumo": _observacao_resumo(wb),
        "extrato_titulo": _celula(wb, "Extrato Bancário", 1, 1),
        "extrato_periodos": _celula(wb, "Extrato Bancário", 2, 1),
        **_pendencias(wb),
    }

    if substituir:
        _backup(cong)
    cong.criar_pastas()
    cong.salvar_config(config)
    cong.salvar_lancamentos(lancamentos)
    cong.salvar_extrato(_tabela(wb, "Extrato Bancário", ["data", "historico", "valor", "situacao", "arquivo"]))
    cong.salvar_comprovantes(_tabela(wb, "Comprovantes WhatsApp", CAMPOS_COMPROVANTE[:-2]))
    cong.salvar_notas(notas)
    destino = cong.entrada("planilha_base") / planilha.name
    if not destino.exists():
        shutil.copy2(planilha, destino)
    return {"lancamentos": len(lancamentos), "meses": len(fontes), "avisos": avisos,
            "revisao_base": config.get("revisao_base")}


def _backup(cong: Congregacao) -> None:
    origem = cong.pasta / "dados"
    if origem.exists():
        destino = cong.cache / f"backup_dados_{datetime.now():%Y%m%d_%H%M%S}"
        shutil.copytree(origem, destino)
        for arquivo in ("transcricoes.yaml", "respostas.yaml", "config.yaml"):
            if (cong.pasta / arquivo).exists():
                shutil.copy2(cong.pasta / arquivo, destino / arquivo)


# --- abas mensais -------------------------------------------------------
def _lancamentos(wb, avisos: list[str]) -> tuple[list[Lancamento], dict]:
    lancamentos: list[Lancamento] = []
    fontes: dict[str, str] = {}
    abas = sorted((mes, ws) for ws in wb.worksheets if (mes := mes_da_aba(ws.title)))
    if ABA_TRANSFERENCIAS in wb.sheetnames:
        # Aba própria das transferências internas: o mês de cada linha vem da coluna "Mês".
        abas.append((None, wb[ABA_TRANSFERENCIAS]))
    for mes_aba, ws in abas:
        mes = mes_aba
        if mes_aba:
            fontes[mes] = str(ws.cell(2, 1).value or "")
        cabecalho = _linha_cabecalho(ws)
        if cabecalho is None:
            avisos.append(f"{ws.title}: cabeçalho 'Data' não encontrado, aba ignorada")
            continue
        for linha in range(cabecalho + 1, ws.max_row + 1):
            valores = [ws.cell(linha, c).value for c in range(1, max(ws.max_column, 8) + 1)]
            if mes_aba is None:
                mes = str(valores[9] or "").strip() if len(valores) > 9 else ""
                valores = valores[:9]
                if not mes:
                    continue  # linhas de total, no fim da aba
            if str(valores[1] or "").strip() == FIM_DOS_LANCAMENTOS:
                break
            if not any(v not in (None, "") for v in valores):
                continue
            tipo = str(valores[1] or "").strip()
            if tipo not in TIPOS:
                avisos.append(f"{ws.title} linha {linha}: sem Entrada/Saída, ignorada: {valores[:8]}")
                continue
            # A coluna "Código" das revisões geradas pelo agente não é dado: os códigos são refeitos a cada importação.
            extras = {f"col{i + 1}": str(v) for i, v in enumerate(valores[8:], start=8)
                      if v not in (None, "") and str(ws.cell(cabecalho, i + 1).value or "") != "Código"}
            try:
                lancamentos.append(Lancamento(
                    id=f"L{len(lancamentos) + 1:05d}", mes=mes, data=para_data(valores[0]), tipo=tipo,
                    descricao=str(valores[2] or "").strip(), valor=dinheiro(valores[3]),
                    forma=str(valores[4] or "Não informado"), comprovante=str(valores[5] or ""),
                    status=normalizar_status(valores[6] or "pendente"),
                    observacao=str(valores[7] or ""), extras=extras))
            except ValueError as erro:
                avisos.append(f"{ws.title} linha {linha}: {erro}")
    # As transferências da aba própria voltam para o fim do seu mês (ordenação estável) e os códigos são refeitos.
    lancamentos.sort(key=lambda l: l.mes)
    for n, l in enumerate(lancamentos, start=1):
        l.id = f"L{n:05d}"
    return lancamentos, fontes


def _linha_cabecalho(ws) -> int | None:
    for linha in range(1, min(ws.max_row, 15) + 1):
        if str(ws.cell(linha, 1).value or "").strip() == "Data":
            return linha
    return None


# --- outras abas --------------------------------------------------------
def _config(wb, planilha: Path, atual: dict) -> dict:
    config = dict(atual)
    regras = dict(config.get("regras", {}))
    titulo = ""
    if "Instruções" in wb.sheetnames:
        ws = wb["Instruções"]
        for linha in range(1, ws.max_row + 1):
            rotulo = str(ws.cell(linha, 2).value or "")
            valor, nota = ws.cell(linha, 3).value, str(ws.cell(linha, 4).value or "")
            if linha <= 2 and rotulo and not titulo:
                titulo = rotulo
            if rotulo.startswith("Repasse Igreja Sede") and valor is not None:
                regras["repasse_sede"] = round(float(valor), 4)
            elif rotulo.startswith("Prebenda Pastoral") and valor is not None:
                regras["prebenda"] = round(float(valor), 4)
            elif rotulo.startswith("Prebenda no saldo") and valor:
                regras["prebenda_no_saldo"] = "quando_paga" if "paga" in str(valor).lower() else "mensal"
            elif rotulo.startswith("Saldo inicial"):
                # Célula com "PREENCHER" = a planilha usou 0 só porque ninguém informou.
                config["saldo_inicial"] = None if "PREENCHER" in nota.upper() else float(valor or 0)
    config["titulo"] = titulo
    if "–" in titulo and not config.get("nome"):
        config["nome"] = titulo.split("–", 1)[1].strip()
    abonos = _abonos_da_planilha(wb, regras.get("abonos_sede") or [])
    if abonos is not None:
        regras["abonos_sede"] = abonos
    config["regras"] = regras
    achado = re.search(r"REV\s*0*(\d+)", planilha.stem, re.IGNORECASE)
    config["revisao_base"] = int(achado.group(1)) if achado else config.get("revisao_base", 0)
    config["planilha_base"] = planilha.name
    return config


def _abonos_da_planilha(wb, atuais: list[dict]) -> list[dict] | None:
    """Meses com repasse à sede 0% (coluna D do resumo) viram períodos de abono.

    Devolve None quando a planilha não traz a porcentagem por mês (revisões antigas):
    aí os abonos do config.yaml continuam valendo."""
    zerados, parciais, achou = [], [], False
    for ws in wb.worksheets:
        mes = mes_da_aba(ws.title)
        if not mes:
            continue
        for linha in range(1, ws.max_row + 1):
            if str(ws.cell(linha, 2).value or "").strip() == "Repasse Igreja Sede":
                valor, nota = ws.cell(linha, 4).value, str(ws.cell(linha, 5).value or "")
                parcial = re.match(r"Abono parcial de (R\$ [\d.,]+)", nota)
                if parcial:
                    achou = True
                    parciais.append({"de": mes, "ate": mes, "valor": float(dinheiro(parcial.group(1)))})
                elif isinstance(valor, (int, float)):
                    achou = True
                    if valor == 0:
                        zerados.append(mes)
                break
    if not achou:
        return None
    motivos = {(a["de"], a["ate"]): a.get("motivo", "") for a in atuais}
    periodos: list[dict] = []
    for mes in sorted(zerados):
        if periodos and proximo_mes(periodos[-1]["ate"]) == mes:
            periodos[-1]["ate"] = mes
        else:
            periodos.append({"de": mes, "ate": mes})
    periodos = sorted(periodos + parciais, key=lambda p: p["de"])
    for p in periodos:
        p["motivo"] = motivos.get((p["de"], p["ate"]), "")
    return periodos


def _situacao(wb) -> str:
    if "Instruções" not in wb.sheetnames:
        return ""
    ws = wb["Instruções"]
    for linha in range(1, ws.max_row):
        if str(ws.cell(linha, 2).value or "").strip() == "SITUAÇÃO ATUAL":
            return str(ws.cell(linha + 1, 2).value or "")
    return ""


def _mes_do_extenso(texto: str) -> str | None:
    achado = re.fullmatch(r"\s*([A-Za-zçÇ]+)/(\d{4})\s*", str(texto or ""))
    if achado and achado.group(1).capitalize() in MESES_PT:
        return f"{achado.group(2)}-{MESES_PT.index(achado.group(1).capitalize()) + 1:02d}"
    return None


def _livro_totais(wb, avisos: list[str]) -> dict:
    if "Conferência Livro" not in wb.sheetnames:
        return {}
    ws, totais = wb["Conferência Livro"], {}
    for linha in range(1, ws.max_row + 1):
        mes = _mes_do_extenso(ws.cell(linha, 1).value)
        if not mes:
            continue
        cel = [ws.cell(linha, c).value for c in range(1, 9)]
        totais[mes] = {"entradas": _numero(cel[2]), "saidas": _numero(cel[5]),
                       "observacao": str(cel[7] or "")}
    return totais


def _numero(valor) -> float | None:
    return round(float(valor), 2) if isinstance(valor, (int, float)) else None


def _observacao_resumo(wb) -> str:
    if "Resumo Geral" not in wb.sheetnames:
        return ""
    ws = wb["Resumo Geral"]
    textos, depois_do_total = [], False
    for linha in range(1, ws.max_row + 1):
        valor = ws.cell(linha, 1).value
        if valor == "TOTAL":
            depois_do_total = True
        elif depois_do_total and isinstance(valor, str):
            textos.append(valor)
    return "\n".join(textos)


def _pendencias(wb) -> dict:
    """Lê a aba Pendências: perguntas viram P1, P2…; as demais seções viram avisos."""
    perguntas, avisos = {}, []
    if "Pendências" not in wb.sheetnames:
        return {"perguntas": perguntas, "avisos": avisos}
    ws, secao = wb["Pendências"], ""
    for linha in range(2, ws.max_row + 1):
        # A REV02 começa na coluna B; as revisões geradas pelo agente, na A.
        inicio = 2 if ws.cell(linha, 1).value in (None, "") else 1
        a, b = ws.cell(linha, inicio).value, ws.cell(linha, inicio + 1).value
        if not a:
            continue
        a = str(a).strip()
        if re.match(r"^\d+\.\s", a):
            secao = a
            if b and not _secao_derivada(secao):
                avisos.append({"secao": secao, "titulo": "", "texto": str(b)})
            continue
        if _secao_derivada(secao):
            continue
        if re.match(r"^\d+ foto\(s\)", str(b or "")):
            continue  # lista antiga de fotos por mês: as pendentes saem do índice de comprovantes
        if "PERGUNTA" in secao.upper():
            perguntas[f"P{len(perguntas) + 1}"] = {"titulo": a, "texto": str(b or "")}
        else:
            avisos.append({"secao": secao, "titulo": a, "texto": str(b or "")})
    return {"perguntas": perguntas, "avisos": avisos}


def _secao_derivada(secao: str) -> bool:
    """Seções recalculadas a cada execução (lançamentos 🔴/🟡 e fotos), não guardadas como texto."""
    return any(chave in secao.upper() for chave in SECOES_DERIVADAS)


def _celula(wb, aba: str, linha: int, coluna: int) -> str:
    return str(wb[aba].cell(linha, coluna).value or "") if aba in wb.sheetnames else ""


def _tabela(wb, aba: str, campos: list[str]) -> list[dict]:
    if aba not in wb.sheetnames:
        return []
    ws = wb[aba]
    cabecalho = _linha_cabecalho(ws)
    if cabecalho is None:
        return []
    linhas = []
    for valores in ws.iter_rows(min_row=cabecalho + 1, max_col=len(campos), values_only=True):
        if not any(v not in (None, "") for v in valores):
            continue
        linha = dict(zip(campos, valores))
        linha["data"] = para_data(linha.get("data"))
        linha["valor"] = dinheiro(linha["valor"]) if isinstance(linha.get("valor"), (int, float)) else None
        linhas.append(linha)
    return linhas
