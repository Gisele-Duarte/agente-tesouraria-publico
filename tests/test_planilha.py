"""A planilha gerada volta para dados sem perder nada, e as revisões nunca se sobrescrevem."""
import openpyxl
import pytest

from tesouraria import planilha
from tesouraria.dados import Congregacao
from tesouraria.importar_planilha import ErroImportacao, importar
from tesouraria.processar import montar_livro, proxima_revisao


def _campos(lancamentos):
    return [(l.mes, l.data, l.tipo, l.descricao, l.valor, l.forma, l.status) for l in lancamentos]


def test_transferencias_internas_ficam_em_aba_propria_e_fora_dos_meses(cong, tmp_path):
    wb = openpyxl.load_workbook(planilha.gerar(montar_livro(cong), tmp_path / "x.xlsx"))
    ws = wb[planilha.ABA_TRANSFERENCIAS]
    assert [ws.cell(r, 3).value for r in (5, 6)] == ["Transferência interna – banco para caixa",
                                                     "Transferência interna – caixa recebe"]
    assert ws.cell(5, 10).value == "2025-01"
    jan = wb["Jan-2025"]
    assert not any(str(jan.cell(r, 3).value or "").startswith("Transferência interna")
                   for r in range(5, jan.max_row + 1))


def test_gerar_e_reimportar_preserva_lancamentos_e_saldos(cong, tmp_path):
    livro = montar_livro(cong)
    arquivo = planilha.gerar(livro, tmp_path / "Livro_Caixa_REV03.xlsx")

    outra = Congregacao(tmp_path / "copia")
    resultado = importar(outra, arquivo)
    copia = montar_livro(outra)

    assert resultado["revisao_base"] == 3
    assert _campos(copia.lancamentos) == _campos(livro.lancamentos)
    assert [r.saldo_total for r in copia.resumos] == [r.saldo_total for r in livro.resumos]
    assert copia.livro_totais["2025-01"]["entradas"] == 1500
    assert set(copia.perguntas) == {"P1", "P2"}


def test_formulas_do_resumo_apontam_para_as_abas_mensais(cong, tmp_path):
    wb = openpyxl.load_workbook(planilha.gerar(montar_livro(cong), tmp_path / "x.xlsx"))
    jan = wb["Jan-2025"]
    rotulos = {jan.cell(r, 2).value: jan.cell(r, 3).value for r in range(1, jan.max_row + 1)}
    assert rotulos["Saldo do mês anterior"] == "=Instruções!$C$10"
    assert '"Transferência interna*"' in rotulos["Entradas (total)"]
    fev = wb["Fev-2025"]
    anterior = next(fev.cell(r, 3).value for r in range(1, fev.max_row + 1)
                    if fev.cell(r, 2).value == "Saldo do mês anterior")
    assert anterior.startswith("='Jan-2025'!$C$")
    assert wb["Resumo Geral"]["K6"].value == "=D6-J6+K5"


def test_importar_nao_sobrescreve_sem_pedir_e_guarda_backup(cong, tmp_path):
    arquivo = planilha.gerar(montar_livro(cong), tmp_path / "REV05.xlsx")
    with pytest.raises(ErroImportacao, match="--substituir"):
        importar(cong, arquivo)
    importar(cong, arquivo, substituir=True)
    assert list(cong.cache.glob("backup_dados_*/lancamentos.csv"))


def test_proxima_revisao_nunca_reaproveita_numero(cong):
    assert proxima_revisao(cong).name == "Livro_Caixa_REV03.xlsx"
    (cong.saida / "Livro_Caixa_REV07.xlsx").write_bytes(b"")
    assert proxima_revisao(cong).name == "Livro_Caixa_REV08.xlsx"


def test_importa_pendencias_no_layout_da_rev02_que_comeca_na_coluna_b(cong, tmp_path):
    arquivo = planilha.gerar(montar_livro(cong), tmp_path / "REV02.xlsx")
    wb = openpyxl.load_workbook(arquivo)
    del wb["Pendências"]
    ws = wb.create_sheet("Pendências")
    for linha in (["PENDÊNCIAS"], [], ["1. EXTRATOS QUE FALTAM"], ["Dez/2024", "Sem extrato"], [],
                  ["2. AGENDA MANUSCRITA", "Fotos"], ["Ago/2025", "21 foto(s)"], [],
                  ["4. PERGUNTAS PARA A LIDERANÇA"], ["Saldo inicial", "Qual era?"], ["Pix de R$ 500", "O que foi?"]):
        ws.append([None, *linha])
    wb.save(arquivo)

    outra = Congregacao(tmp_path / "rev02")
    importar(outra, arquivo)
    notas = outra.notas()
    assert [p["titulo"] for p in notas["perguntas"].values()] == ["Saldo inicial", "Pix de R$ 500"]
    assert notas["avisos"] == [{"secao": "1. EXTRATOS QUE FALTAM", "titulo": "Dez/2024", "texto": "Sem extrato"}]


def test_revisao_marca_novos_e_alterados_e_lista_as_alteracoes(cong, tmp_path):
    from conftest import escrever
    escrever(cong, "respostas.yaml", {
        "base_revisao": 2,
        "lancamentos": {"L00007": {"resposta": "foram 80", "valor": 80, "status": "confirmado"},
                        "L00010": {"resposta": "lançado em dobro", "remover": True}},
        "novos": [{"resposta": "faltou", "data": "2025-02-10", "tipo": "Entrada", "descricao": "Oferta – culto",
                   "valor": 5, "forma": "Dinheiro"}],
    })
    livro = montar_livro(cong)
    assert {a["id"]: a["tipo"] for a in livro.alteracoes} == {"L00007": "Alterado", "R0001": "Novo",
                                                             "L00010": "Removido"}
    alterado = next(a for a in livro.alteracoes if a["id"] == "L00007")
    assert "Valor: (vazio) → 80,00" in alterado["mudou"] and "foram 80" in alterado["motivo"]
    wb = openpyxl.load_workbook(planilha.gerar(livro, tmp_path / "Livro_Caixa_REV03.xlsx"))
    aba = wb["Alterações REV03"]
    assert [aba.cell(r, 1).value for r in range(5, 8)] == ["Alterado", "Removido", "Novo"]  # por data
    fev = wb["Fev-2025"]
    linha = next(r for r in range(5, fev.max_row + 1) if fev.cell(r, 9).value == "R0001")
    assert fev.cell(linha, 1).fill.fgColor.rgb.endswith("DDEBF7")
    # reimportar não transforma a coluna Código em dado
    outra = Congregacao(tmp_path / "rev03")
    importar(outra, tmp_path / "Livro_Caixa_REV03.xlsx")
    assert all(not l.extras for l in outra.lancamentos())


def test_abono_vai_para_a_planilha_e_volta_na_importacao(cong, tmp_path):
    config = cong.config()
    config["regras"]["abonos_sede"] = [{"de": "2025-02", "ate": "2025-02", "motivo": "abonado"}]
    cong.salvar_config(config)
    livro = montar_livro(cong)
    arquivo = planilha.gerar(livro, tmp_path / "Livro_Caixa_REV03.xlsx")
    fev = openpyxl.load_workbook(arquivo)["Fev-2025"]
    linha = next(r for r in range(1, fev.max_row + 1) if fev.cell(r, 2).value == "Repasse Igreja Sede")
    assert fev.cell(linha, 4).value == 0 and fev.cell(linha, 3).value == f"=C{linha - 1}*D{linha}"
    outra = Congregacao(tmp_path / "rev03")
    importar(outra, arquivo)
    assert outra.config()["regras"]["abonos_sede"] == [{"de": "2025-02", "ate": "2025-02", "motivo": ""}]
    assert [r.saldo_total for r in montar_livro(outra).resumos] == [r.saldo_total for r in livro.resumos]


def test_regras_novas_vao_para_a_planilha_e_voltam(cong, tmp_path):
    config = cong.config()
    config["regras"]["prebenda_no_saldo"] = "quando_paga"
    config["regras"]["abonos_sede"] = [{"de": "2025-01", "ate": "2025-01", "valor": 100, "motivo": "parcial"},
                                       {"de": "2025-02", "ate": "2025-02", "motivo": "abonado"}]
    cong.salvar_config(config)
    livro = montar_livro(cong)
    arquivo = planilha.gerar(livro, tmp_path / "Livro_Caixa_REV03.xlsx")
    outra = Congregacao(tmp_path / "rev03")
    importar(outra, arquivo)
    regras = outra.config()["regras"]
    assert regras["prebenda_no_saldo"] == "quando_paga"
    assert [(a["de"], a.get("valor")) for a in regras["abonos_sede"]] == [("2025-01", 100.0), ("2025-02", None)]
    assert [r.saldo_total for r in montar_livro(outra).resumos] == [r.saldo_total for r in livro.resumos]


def test_mes_abonado_mostra_os_20_por_cento_como_abonado_sem_mudar_o_saldo(cong, tmp_path):
    sem_abono = [r.saldo_total for r in montar_livro(cong).resumos]
    config = cong.config()
    config["regras"]["abonos_sede"] = [{"de": "2025-01", "ate": "2025-01", "motivo": "abonado"}]
    cong.salvar_config(config)
    livro = montar_livro(cong)
    arquivo = planilha.gerar(livro, tmp_path / "Livro_Caixa_REV03.xlsx")
    jan = openpyxl.load_workbook(arquivo)["Jan-2025"]
    linha = next(r for r in range(1, jan.max_row + 1) if jan.cell(r, 2).value == "Repasse Igreja Sede")
    assert jan.cell(linha, 6).value == f"=C{linha - 1}*Instruções!$C$6"
    assert jan.cell(linha, 7).value == planilha.ROTULO_ABONADO
    # o valor abonado é só informativo: o repasse do mês continua 0 e a igreja local fica com mais
    assert jan.cell(linha, 4).value == 0
    assert [r.saldo_total for r in livro.resumos][0] > sem_abono[0]


def test_aba_gestao_anterior_sai_na_planilha_e_nao_entra_no_livro(cong, tmp_path):
    from tesouraria.dados import gravar_yaml
    gravar_yaml(cong.pasta / "gestao_anterior.yaml", {
        "titulo": "GESTÃO ANTERIOR – teste",
        "lancamentos": [{"data": "2024-12-01", "tipo": "Entrada", "descricao": "Oferta", "valor": 10.0,
                         "fonte": "foto teste", "observacao": ""}],
        "resumo": [{"rotulo": "Saldo final anotado", "valor": -5.0, "observacao": "teste"}]})
    antes = [r.saldo_total for r in montar_livro(cong).resumos]
    livro = montar_livro(cong)
    arquivo = planilha.gerar(livro, tmp_path / "Livro_Caixa_REV03.xlsx")
    ws = openpyxl.load_workbook(arquivo)[planilha.ABA_GESTAO_ANTERIOR]
    assert ws["A1"].value == "GESTÃO ANTERIOR – teste" and ws["C5"].value == "Oferta"
    assert [r.saldo_total for r in livro.resumos] == antes
    outra = Congregacao(tmp_path / "rev03")
    importar(outra, arquivo)
    assert _campos(montar_livro(outra).lancamentos) == _campos(livro.lancamentos)


def test_valores_na_sede_lista_cartao_e_pix_sem_extrato_sem_mudar_o_saldo(cong, tmp_path):
    from datetime import date
    from decimal import Decimal
    from tesouraria.modelo import Lancamento
    from tesouraria.regras import MOTIVO_CARTAO, MOTIVO_PIX_SEM_EXTRATO, Regras, motivo_na_sede
    regras = Regras(na_sede_desde="2025-02")
    def l(mes, forma, desc, comp="Livro caixa"):
        return Lancamento(id="X", mes=mes, data=date(2025, int(mes[-2:]), 1), tipo="Entrada", descricao=desc,
                          valor=Decimal("10"), forma=forma, comprovante=comp, status="🟢 CONFIRMADO")
    assert motivo_na_sede(l("2025-02", "Cartão", "Oferta – culto (cartão)"), regras) == MOTIVO_CARTAO
    assert motivo_na_sede(l("2025-02", "Pix", "Dízimo – Fulano"), regras) == MOTIVO_PIX_SEM_EXTRATO
    assert motivo_na_sede(l("2025-02", "Pix", "Dízimo – Fulano", "Extrato 02/02"), regras) is None
    assert motivo_na_sede(l("2025-02", "Dinheiro", "Oferta – culto (dinheiro)"), regras) is None
    assert motivo_na_sede(l("2025-01", "Cartão", "Oferta – culto (cartão)"), regras) is None

    antes = [r.saldo_total for r in montar_livro(cong).resumos]
    config = cong.config()
    config["regras"]["valores_na_sede"] = {"desde": "2025-01", "motivo": "teste"}
    cong.salvar_config(config)
    livro = montar_livro(cong)
    assert [r.saldo_total for r in livro.resumos] == antes
    arquivo = planilha.gerar(livro, tmp_path / "Livro_Caixa_REV03.xlsx")
    wb = openpyxl.load_workbook(arquivo)
    ws = wb[planilha.ABA_NA_SEDE]
    codigos = {ws.cell(r, 8).value for r in range(5, ws.max_row + 1)} - {None}
    assert {"L00001", "L00002"} <= codigos          # Pix do livro sem extrato, em jan/2025
    resumo = wb["Resumo Geral"]
    assert resumo.cell(4, 17).value == "Na Sede (acum.)" and resumo.cell(5, 18).value == "=I5-Q5"
    outra = Congregacao(tmp_path / "rev03")
    importar(outra, arquivo)
    assert _campos(montar_livro(outra).lancamentos) == _campos(livro.lancamentos)
